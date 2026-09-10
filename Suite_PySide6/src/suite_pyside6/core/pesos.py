from __future__ import annotations

import os
import json
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Callable, Literal
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


TARGET_SHEET_NAME = "Hoja1"
OOXML_EXTENSIONS = {".xlsx", ".xlsm"}
LEGACY_EXCEL_EXTENSIONS = {".xls"}
SUPPORTED_EXTENSIONS = OOXML_EXTENSIONS | LEGACY_EXCEL_EXTENSIONS
OLD_EXCEL_EXTENSIONS = {".xlsb"}
WORKBOOK_XML = "xl/workbook.xml"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

VaciadoType = Literal["ninguno", "normal", "completo"]
VALID_VACIADO_TYPES = frozenset(("ninguno", "normal", "completo"))


ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
ET.register_namespace("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
ET.register_namespace("x15", "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main")
ET.register_namespace("xr", "http://schemas.microsoft.com/office/spreadsheetml/2014/revision")
ET.register_namespace("xr6", "http://schemas.microsoft.com/office/spreadsheetml/2016/revision6")
ET.register_namespace("xr10", "http://schemas.microsoft.com/office/spreadsheetml/2016/revision10")


@dataclass(frozen=True)
class SheetRename:
    path: Path
    success: bool
    before: str = ""
    after: str = TARGET_SHEET_NAME
    changed: bool = False
    message: str = ""
    vaciado: VaciadoType = "ninguno"
    adjusted_sheets: int = 0
    adjusted_weights: int = 0


@dataclass(frozen=True)
class PesosProgress:
    """A unit of real work completed while processing a group of workbooks."""

    completed: int
    total: int
    message: str


@dataclass(frozen=True)
class _WeightSheetPlan:
    title: str
    header_row: int
    weight_column: int
    rows: tuple[int, ...]


@dataclass(frozen=True)
class _FilePlan:
    path: Path
    vaciado: VaciadoType
    weight_sheets: tuple[_WeightSheetPlan, ...] = ()
    legacy_weight_count: int | None = None

    @property
    def weight_count(self) -> int:
        return self.legacy_weight_count if self.legacy_weight_count is not None else sum(len(sheet.rows) for sheet in self.weight_sheets)


@dataclass
class PesosResult:
    selected_files: list[Path] = field(default_factory=list)
    results: list[SheetRename] = field(default_factory=list)
    ignored_files: list[Path] = field(default_factory=list)

    @property
    def processed_count(self) -> int:
        return sum(1 for item in self.results if item.success and item.changed)

    @property
    def ok_count(self) -> int:
        return sum(1 for item in self.results if item.success)

    @property
    def unchanged_count(self) -> int:
        return sum(1 for item in self.results if item.success and not item.changed)

    @property
    def error_count(self) -> int:
        return sum(1 for item in self.results if not item.success) + len(self.ignored_files)

    def summary(self) -> str:
        return (
            f"Archivos seleccionados: {len(self.selected_files)} | "
            f"Renombrados: {self.processed_count} | "
            f"Ya estaban correctos: {self.unchanged_count} | "
            f"Errores/ignorados: {self.error_count}"
        )

    def log_text(self) -> str:
        if not self.selected_files:
            return "Selecciona archivos Excel para empezar."

        lines = ["Resultado del proceso:"]
        for item in self.results:
            adjustment = (
                f" Vaciado {item.vaciado}: {item.adjusted_weights} peso(s) en {item.adjusted_sheets} hoja(s)."
                if item.vaciado != "ninguno" and item.success
                else ""
            )
            if item.success and item.changed:
                lines.append(f"- OK {item.path.name}: '{item.before}' -> '{item.after}'.{adjustment}")
            elif item.success:
                lines.append(f"- OK {item.path.name}: la hoja ya se llamaba '{TARGET_SHEET_NAME}'.{adjustment}")
            else:
                lines.append(f"- ERROR {item.path.name}: {item.message}")

        for path in self.ignored_files:
            ext = path.suffix.lower()
            if ext in OLD_EXCEL_EXTENSIONS:
                lines.append(f"- Ignorado {path.name}: formato Excel antiguo/no soportado.")
            else:
                lines.append(f"- Ignorado {path.name}: no es un Excel .xlsx/.xlsm/.xls.")
        return "\n".join(lines)


def calcular_peso_vaciado(peso_bruto: object, tipo_vaciado: VaciadoType) -> str:
    """Apply the business rule and return the Excel representation (one decimal).

    Decimal plus ROUND_HALF_UP is intentional: Python's default rounding is
    banker's rounding and differs from Excel's ROUND for values ending in 5.
    """
    value = _calcular_peso_vaciado_decimal(peso_bruto, tipo_vaciado)
    return _format_decimal(value)


def _calcular_peso_vaciado_decimal(peso_bruto: object, tipo_vaciado: VaciadoType) -> Decimal:
    """Return the numeric result without passing through a binary float."""
    if tipo_vaciado not in VALID_VACIADO_TYPES:
        raise ValueError(f"tipo de vaciado no valido: {tipo_vaciado}")
    value = _as_decimal(peso_bruto)
    if tipo_vaciado == "ninguno":
        return value
    adjusted = value - (value * Decimal("0.011"))
    if tipo_vaciado == "completo":
        adjusted -= Decimal("2.9")
    # One decimal is an explicit rule of Vaciado, not an intermediate rounding.
    return adjusted.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def process_pesos_files(
    paths: list[Path],
    vaciados: dict[Path, VaciadoType] | None = None,
    progress: Callable[[PesosProgress], None] | None = None,
) -> PesosResult:
    """Rename the current target sheet and optionally adjust lote weights.

    Every workbook is validated before its original file is replaced.  XLSX
    workbooks are handled in memory through openpyxl; legacy XLS keeps the
    suite's existing Excel-compatible route so its original format is retained.
    """
    result = PesosResult(selected_files=list(paths))
    selection = vaciados or {}
    plans: list[_FilePlan] = []

    for path in paths:
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            result.ignored_files.append(path)
            continue
        try:
            vaciado = _normalise_vaciado(selection.get(path, "ninguno"))
            plans.append(_build_file_plan(path, vaciado))
        except Exception as exc:
            result.results.append(SheetRename(path=path, success=False, message=str(exc)))

    # Validation is a completed, measured phase.  Remaining units are one
    # rename plus one save per valid file, and one unit for each adjusted cell.
    total = len(paths) + sum(2 + plan.weight_count for plan in plans)
    completed = len(paths)
    if result.error_count:
        # Reserve the final unit so an error can never be presented as 100%.
        total = max(total, completed + 1)
    _emit_progress(progress, completed, total, "Validando archivos…")

    # A validation failure is reported before any source workbook is changed.
    # It is safer to ask the operator to correct the batch than to leave a
    # mixed batch where only the first files were written successfully.
    if result.error_count:
        _emit_progress(progress, completed, total, "Validación detenida por errores.")
        return result

    for plan in plans:
        _emit_progress(progress, completed, total, f"Renombrando hoja de {plan.path.name}…")
        try:
            item = _process_file_plan(plan, progress, completed, total)
            completed += 1 + plan.weight_count + 1
            result.results.append(item)
            _emit_progress(progress, completed, total, f"Guardado {plan.path.name}.")
        except Exception as exc:
            result.results.append(
                SheetRename(path=plan.path, success=False, message=str(exc), vaciado=plan.vaciado)
            )
            # Do not claim completion after an error: callers keep the actual value.
            _emit_progress(progress, completed, total, f"Error al procesar {plan.path.name}.")
            break
    else:
        _emit_progress(progress, total, total, "Proceso completado.")
    return result


def _normalise_vaciado(value: object) -> VaciadoType:
    normalised = str(value).strip().casefold()
    if normalised not in VALID_VACIADO_TYPES:
        raise ValueError(f"tipo de vaciado no valido: {value}")
    return normalised  # type: ignore[return-value]


def _as_decimal(value: object) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError("no es un numero valido")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "")
    if not text:
        raise ValueError("no es un numero valido")
    # Accept the usual Excel imports without changing their textual contents.
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("no es un numero valido") from exc


def _format_decimal(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP), ".1f")


def _normalise_header(value: object) -> str:
    return "".join(str(value or "").split()).casefold()


def _build_file_plan(path: Path, vaciado: VaciadoType) -> _FilePlan:
    if vaciado == "ninguno":
        return _FilePlan(path=path, vaciado=vaciado)
    if path.suffix.lower() in LEGACY_EXCEL_EXTENSIONS:
        return _FilePlan(path=path, vaciado=vaciado, legacy_weight_count=_inspect_legacy_xls(path, vaciado))
    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=keep_vba)
    try:
        plans = tuple(_weight_sheet_plan(sheet) for sheet in workbook.worksheets if _sheet_has_weight_header(sheet))
        if not plans:
            raise ValueError(f"{path.name}: no se encontro el encabezado pesoBruto en ninguna hoja")
        return _FilePlan(path=path, vaciado=vaciado, weight_sheets=plans)
    finally:
        workbook.close()


def _sheet_has_weight_header(sheet: object) -> bool:
    return any(
        _normalise_header(cell.value) == "pesobruto"
        for row in sheet.iter_rows()  # type: ignore[attr-defined]
        for cell in row
    )


def _weight_sheet_plan(sheet: object) -> _WeightSheetPlan:
    matches: list[tuple[int, int]] = []
    for row in sheet.iter_rows():  # type: ignore[attr-defined]
        for cell in row:
            if _normalise_header(cell.value) == "pesobruto":
                matches.append((cell.row, cell.column))
    if len(matches) != 1:
        raise ValueError(f"hoja {sheet.title}: encabezado pesoBruto ambiguo")
    header_row, weight_column = matches[0]
    rows: list[int] = []
    for row_index in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(row_index, column).value for column in range(1, sheet.max_column + 1)]
        if not any(value not in (None, "") for value in values):
            continue
        weight = sheet.cell(row_index, weight_column).value
        if weight in (None, ""):
            continue
        try:
            _as_decimal(weight)
        except ValueError as exc:
            raise ValueError(f"hoja {sheet.title}, fila {row_index}, pesoBruto: {exc}") from exc
        rows.append(row_index)
    return _WeightSheetPlan(sheet.title, header_row, weight_column, tuple(rows))


def _process_file_plan(
    plan: _FilePlan,
    progress: Callable[[PesosProgress], None] | None,
    completed: int,
    total: int,
) -> SheetRename:
    if plan.path.suffix.lower() in LEGACY_EXCEL_EXTENSIONS:
        _emit_progress(progress, completed + 1, total, f"Ajustando pesos de {plan.path.name}…")
        return _process_legacy_xls(plan, progress, completed, total)
    return _process_ooxml(plan, progress, completed, total)


def _process_ooxml(
    plan: _FilePlan,
    progress: Callable[[PesosProgress], None] | None,
    completed: int,
    total: int,
) -> SheetRename:
    if plan.vaciado == "ninguno":
        return _rename_ooxml_atomically(plan.path)
    keep_vba = plan.path.suffix.lower() == ".xlsm"
    workbook = load_workbook(plan.path, keep_vba=keep_vba, data_only=False)
    temp_path: Path | None = None
    try:
        target_sheet = _first_visible_openpyxl_sheet(workbook.worksheets)
        before = target_sheet.title
        changed = before != TARGET_SHEET_NAME
        if changed:
            if any(sheet is not target_sheet and sheet.title.casefold() == TARGET_SHEET_NAME.casefold() for sheet in workbook.worksheets):
                raise ValueError("ya existe otra hoja llamada Hoja1; no se cambia nada para evitar nombres duplicados")

        adjusted_sheets = 0
        adjusted_weights = 0
        if plan.vaciado != "ninguno":
            for sheet_plan in plan.weight_sheets:
                sheet = workbook[sheet_plan.title]
                adjusted_sheets += 1
                for position, row in enumerate(sheet_plan.rows, start=1):
                    cell = sheet.cell(row, sheet_plan.weight_column)
                    cell.value = _calcular_peso_vaciado_decimal(cell.value, plan.vaciado)
                    cell.number_format = "0.0"
                    adjusted_weights += 1
                    _emit_progress(
                        progress,
                        completed + 1 + adjusted_weights,
                        total,
                        f"Procesando peso {position} de {len(sheet_plan.rows)} en {plan.path.name}…",
                    )
        if changed:
            target_sheet.title = TARGET_SHEET_NAME
        temp_path = _temporary_path_for(plan.path)
        workbook.save(temp_path)
        os.replace(temp_path, plan.path)
        temp_path = None
        return SheetRename(
            path=plan.path,
            success=True,
            before=before,
            changed=changed,
            vaciado=plan.vaciado,
            adjusted_sheets=adjusted_sheets,
            adjusted_weights=adjusted_weights,
        )
    finally:
        workbook.close()
        if temp_path is not None and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _rename_ooxml_atomically(path: Path) -> SheetRename:
    """Keep the original XML-only rename route for files without vaciado."""
    temp_path = _temporary_path_for(path)
    try:
        shutil.copy2(path, temp_path)
        temporary_result = rename_first_visible_sheet(temp_path)
        if not temporary_result.success:
            raise ValueError(temporary_result.message or "no se pudo renombrar la hoja")
        os.replace(temp_path, path)
        return SheetRename(
            path=path,
            success=True,
            before=temporary_result.before,
            changed=temporary_result.changed,
        )
    finally:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _first_visible_openpyxl_sheet(sheets: list[object]) -> object:
    for sheet in sheets:
        if sheet.sheet_state == "visible":
            return sheet
    if not sheets:
        raise ValueError("el libro no contiene hojas")
    return sheets[0]


def _temporary_path_for(path: Path) -> Path:
    with tempfile.NamedTemporaryFile(delete=False, dir=path.parent, suffix=path.suffix) as temp_file:
        return Path(temp_file.name)


def _emit_progress(
    callback: Callable[[PesosProgress], None] | None,
    completed: int,
    total: int,
    message: str,
) -> None:
    if callback is not None:
        callback(PesosProgress(completed=max(0, completed), total=max(1, total), message=message))


def _inspect_legacy_xls(path: Path, vaciado: VaciadoType) -> int:
    """Validate a legacy workbook and count its weight cells before work starts."""
    script_path: Path | None = None
    try:
        script_path = _write_legacy_adjustment_script()
        completed = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script_path),
                str(path),
                TARGET_SHEET_NAME,
                vaciado,
                "preview",
            ],
            check=False,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=180,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        data = _json_last_line(completed.stdout)
        if completed.returncode != 0 or not data.get("success"):
            raise ValueError(str(data.get("message") or completed.stderr.strip() or "no se pudo validar el archivo .xls"))
        return int(data.get("plannedWeights") or 0)
    except FileNotFoundError as exc:
        raise ValueError("no se encontro PowerShell para procesar archivos .xls") from exc
    except subprocess.TimeoutExpired as exc:
        raise ValueError("Excel no respondio al validar el archivo .xls") from exc
    finally:
        if script_path is not None:
            script_path.unlink(missing_ok=True)


def _process_legacy_xls(
    plan: _FilePlan,
    progress: Callable[[PesosProgress], None] | None,
    completed_units: int,
    total_units: int,
) -> SheetRename:
    """Use a temporary copy so a failed XLS operation never touches source."""
    temp_path = _temporary_path_for(plan.path)
    script_path: Path | None = None
    try:
        shutil.copy2(plan.path, temp_path)
        script_path = _write_legacy_adjustment_script()
        process = subprocess.Popen(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script_path),
                str(temp_path),
                TARGET_SHEET_NAME,
                plan.vaciado,
            ],
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        output_lines: list[str] = []
        assert process.stdout is not None
        for line in process.stdout:
            line = line.rstrip("\r\n")
            if line.startswith("PROGRESS|"):
                _, processed, planned = line.split("|", maxsplit=2)
                _emit_progress(
                    progress,
                    completed_units + 1 + int(processed),
                    total_units,
                    f"Procesando peso {processed} de {planned} en {plan.path.name}…",
                )
            else:
                output_lines.append(line)
        try:
            return_code = process.wait(timeout=180)
        except subprocess.TimeoutExpired as exc:
            process.kill()
            raise exc
        output = "\n".join(output_lines)
        data = _json_last_line(output)
        if return_code != 0 or not data.get("success"):
            raise ValueError(str(data.get("message") or output.strip() or "no se pudo procesar el archivo .xls"))
        os.replace(temp_path, plan.path)
        return SheetRename(
            path=plan.path,
            success=True,
            before=str(data.get("before") or ""),
            changed=bool(data.get("changed")),
            vaciado=plan.vaciado,
            adjusted_sheets=int(data.get("adjustedSheets") or 0),
            adjusted_weights=int(data.get("adjustedWeights") or 0),
        )
    except FileNotFoundError as exc:
        raise ValueError("no se encontro PowerShell para procesar archivos .xls") from exc
    except subprocess.TimeoutExpired as exc:
        raise ValueError("Excel no respondio al procesar el archivo .xls") from exc
    finally:
        if script_path is not None:
            script_path.unlink(missing_ok=True)
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _json_last_line(output: str) -> dict[str, object]:
    try:
        data = json.loads(output.strip().splitlines()[-1]) if output.strip() else {}
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def _write_legacy_adjustment_script() -> Path:
    # This remains deliberately self-contained because pyproject intentionally
    # has no XLS writer dependency.  COM is only the legacy fallback; XLSX/XLSM
    # processing above does not require Excel to be installed.
    script = r'''
param([string]$Path, [string]$TargetName, [string]$Vaciado, [string]$Preview = '')
$ErrorActionPreference = "Stop"
$excel = $null
$workbook = $null
$isPreview = $Preview -eq 'preview'
function Header-Key([object]$Value) {
    return (([string]$Value -replace '\s', '').ToLowerInvariant())
}
function Decimal-Value([object]$Value) {
    if ($null -eq $Value -or $Value -is [bool]) { throw "no es un numero valido" }
    $invariant = [Globalization.CultureInfo]::InvariantCulture
    $styles = [Globalization.NumberStyles]::AllowLeadingWhite -bor [Globalization.NumberStyles]::AllowTrailingWhite -bor [Globalization.NumberStyles]::AllowLeadingSign -bor [Globalization.NumberStyles]::AllowDecimalPoint -bor [Globalization.NumberStyles]::AllowExponent
    # Value2 returns native numeric COM values.  Converting those to a string
    # with the current culture first turns 143.7 into "143,7" in Spanish;
    # invariant Number parsing then reads that comma as a thousands separator.
    # Format numerics invariantly before parsing so no decimal is displaced.
    if ($Value -is [byte] -or $Value -is [sbyte] -or $Value -is [int16] -or $Value -is [uint16] -or $Value -is [int32] -or $Value -is [uint32] -or $Value -is [int64] -or $Value -is [uint64] -or $Value -is [single] -or $Value -is [double] -or $Value -is [decimal]) {
        $text = [Convert]::ToString($Value, $invariant)
        return [decimal]::Parse($text, $styles, $invariant)
    }
    $text = ([string]$Value).Trim()
    if ([string]::IsNullOrWhiteSpace($text)) { throw "no es un numero valido" }
    $text = $text -replace '\s', ''
    if ($text.Contains(',') -and $text.Contains('.')) {
        if ($text.LastIndexOf(',') -gt $text.LastIndexOf('.')) {
            $text = $text.Replace('.', '').Replace(',', '.')
        } else {
            $text = $text.Replace(',', '')
        }
    } elseif ($text.Contains(',')) {
        $text = $text.Replace(',', '.')
    }
    try { return [decimal]::Parse($text, $styles, $invariant) }
    catch { throw "no es un numero valido" }
}
try {
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $workbook = $excel.Workbooks.Open($Path)
    $targetSheet = $null
    foreach ($sheet in @($workbook.Worksheets)) {
        if ($sheet.Visible -eq -1) { $targetSheet = $sheet; break }
    }
    if ($null -eq $targetSheet) { $targetSheet = $workbook.Worksheets.Item(1) }
    $previous = [string]$targetSheet.Name
    $changed = $false
    if ($previous -ne $TargetName) {
        foreach ($sheet in @($workbook.Worksheets)) {
            if ($sheet.Name -ieq $TargetName -and $sheet.Index -ne $targetSheet.Index) {
                throw "ya existe otra hoja llamada $TargetName; no se cambia nada para evitar nombres duplicados"
            }
        }
        if (-not $isPreview) {
            $targetSheet.Name = $TargetName
            $changed = $true
        }
    }
    $adjustedSheets = 0
    $adjustedWeights = 0
    $weightCells = @()
    if ($Vaciado -ne 'ninguno') {
        foreach ($sheet in @($workbook.Worksheets)) {
            $used = $sheet.UsedRange
            $headerHits = @()
            for ($r = 1; $r -le $used.Rows.Count; $r++) {
                for ($c = 1; $c -le $used.Columns.Count; $c++) {
                    $cell = $used.Cells.Item($r, $c)
                    if ((Header-Key $cell.Value2) -eq 'pesobruto') {
                        $headerHits += ,@($cell.Row, $cell.Column)
                    }
                }
            }
            if ($headerHits.Count -eq 0) { continue }
            if ($headerHits.Count -ne 1) { throw "hoja $($sheet.Name): encabezado pesoBruto ambiguo" }
            $header = $headerHits[0]
            $headerRow = [int]$header[0]
            $weightColumn = [int]$header[1]
            $adjustedSheets++
            $lastRow = $used.Row + $used.Rows.Count - 1
            for ($r = $headerRow + 1; $r -le $lastRow; $r++) {
                $hasData = $false
                for ($c = $used.Column; $c -lt ($used.Column + $used.Columns.Count); $c++) {
                    $candidate = $sheet.Cells.Item($r, $c).Value2
                    if ($null -ne $candidate -and [string]$candidate -ne '') { $hasData = $true; break }
                }
                if (-not $hasData) { continue }
                $weightCell = $sheet.Cells.Item($r, $weightColumn)
                if ($null -eq $weightCell.Value2 -or [string]$weightCell.Value2 -eq '') { continue }
                try { $weight = Decimal-Value $weightCell.Value2 }
                catch { throw "hoja $($sheet.Name), fila $r, pesoBruto: $($_.Exception.Message)" }
                $weightCells += [pscustomobject]@{Cell=$weightCell; Weight=$weight}
            }
        }
        if ($adjustedSheets -eq 0) { throw "no se encontro el encabezado pesoBruto en ninguna hoja" }
    }
    $plannedWeights = $weightCells.Count
    if (-not $isPreview) {
        foreach ($entry in $weightCells) {
            $adjusted = $entry.Weight - ($entry.Weight * [decimal]0.011)
            if ($Vaciado -eq 'completo') { $adjusted -= [decimal]2.9 }
            $adjusted = [Math]::Round($adjusted, 1, [MidpointRounding]::AwayFromZero)
            $entry.Cell.Value2 = $adjusted
            $entry.Cell.NumberFormat = '0.0'
            $adjustedWeights++
            Write-Output "PROGRESS|$adjustedWeights|$plannedWeights"
        }
        $workbook.Save()
    }
    [pscustomobject]@{success=$true; before=$previous; changed=$changed; adjustedSheets=$adjustedSheets; adjustedWeights=$adjustedWeights; plannedWeights=$plannedWeights} | ConvertTo-Json -Compress
} catch {
    [pscustomobject]@{success=$false; message=$_.Exception.Message} | ConvertTo-Json -Compress
    exit 2
} finally {
    if ($workbook -ne $null) { $workbook.Close($false) }
    if ($excel -ne $null) { $excel.Quit(); [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
'''
    with tempfile.NamedTemporaryFile("w", delete=False, suffix=".ps1", encoding="utf-8") as temp_file:
        temp_file.write(script)
        return Path(temp_file.name)


def rename_first_visible_sheet(path: Path) -> SheetRename:
    if path.suffix.lower() in LEGACY_EXCEL_EXTENSIONS:
        return _rename_legacy_xls_with_excel(path)
    try:
        with ZipFile(path, "r") as source:
            if WORKBOOK_XML not in source.namelist():
                raise ValueError("no se encontro xl/workbook.xml")
            workbook_xml = source.read(WORKBOOK_XML)
        new_workbook_xml, previous_name, changed = _renamed_workbook_xml(workbook_xml)
        if not changed:
            return SheetRename(path=path, success=True, before=previous_name, changed=False)
        _replace_workbook_xml(path, new_workbook_xml)
        return SheetRename(path=path, success=True, before=previous_name, changed=True)
    except BadZipFile:
        return SheetRename(path=path, success=False, message="el archivo no parece un Excel XLSX/XLSM valido")
    except PermissionError:
        return SheetRename(path=path, success=False, message="no se pudo escribir; cierra el archivo en Excel")
    except Exception as exc:
        return SheetRename(path=path, success=False, message=str(exc))


def _renamed_workbook_xml(workbook_xml: bytes) -> tuple[bytes, str, bool]:
    root = ET.fromstring(workbook_xml)
    sheets = root.find(f"{{{MAIN_NS}}}sheets")
    if sheets is None:
        raise ValueError("el libro no contiene hojas")

    sheet_nodes = list(sheets.findall(f"{{{MAIN_NS}}}sheet"))
    if not sheet_nodes:
        raise ValueError("el libro no contiene hojas")

    target_sheet = _first_visible_sheet(sheet_nodes)
    previous_name = target_sheet.attrib.get("name", "")
    if previous_name == TARGET_SHEET_NAME:
        return workbook_xml, previous_name, False

    for sheet in sheet_nodes:
        if sheet is not target_sheet and sheet.attrib.get("name", "").casefold() == TARGET_SHEET_NAME.casefold():
            raise ValueError("ya existe otra hoja llamada Hoja1; no se cambia nada para evitar nombres duplicados")

    target_sheet.set("name", TARGET_SHEET_NAME)
    updated = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return updated, previous_name, True


def _first_visible_sheet(sheet_nodes: list[ET.Element]) -> ET.Element:
    for sheet in sheet_nodes:
        if sheet.attrib.get("state", "visible") == "visible":
            return sheet
    return sheet_nodes[0]


def _replace_workbook_xml(path: Path, workbook_xml: bytes) -> None:
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, dir=path.parent, suffix=path.suffix) as temp_file:
            temp_path = Path(temp_file.name)

        with ZipFile(path, "r") as source, ZipFile(temp_path, "w") as target:
            workbook_written = False
            for info in source.infolist():
                if info.filename == WORKBOOK_XML:
                    if not workbook_written:
                        target.writestr(info, workbook_xml)
                        workbook_written = True
                    continue
                target.writestr(info, source.read(info.filename))

        os.replace(temp_path, path)
    except Exception:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise


def _rename_legacy_xls_with_excel(path: Path) -> SheetRename:
    script = r"""
param([string]$Path, [string]$TargetName)
$ErrorActionPreference = "Stop"
$excel = $null
$workbook = $null
try {
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $workbook = $excel.Workbooks.Open($Path)
    $targetSheet = $null
    foreach ($sheet in @($workbook.Worksheets)) {
        if ($sheet.Visible -eq -1) {
            $targetSheet = $sheet
            break
        }
    }
    if ($null -eq $targetSheet) {
        $targetSheet = $workbook.Worksheets.Item(1)
    }
    $previous = [string]$targetSheet.Name
    $changed = $false
    if ($previous -ne $TargetName) {
        foreach ($sheet in @($workbook.Worksheets)) {
            if ($sheet.Name -ieq $TargetName -and $sheet.Index -ne $targetSheet.Index) {
                throw "ya existe otra hoja llamada $TargetName; no se cambia nada para evitar nombres duplicados"
            }
        }
        $targetSheet.Name = $TargetName
        $workbook.Save()
        $changed = $true
    }
    [pscustomobject]@{success=$true; before=$previous; changed=$changed} | ConvertTo-Json -Compress
} catch {
    [pscustomobject]@{success=$false; message=$_.Exception.Message} | ConvertTo-Json -Compress
    exit 2
} finally {
    if ($workbook -ne $null) {
        $workbook.Close($false)
    }
    if ($excel -ne $null) {
        $excel.Quit()
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel)
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""
    script_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile("w", delete=False, suffix=".ps1", encoding="utf-8") as temp_file:
            temp_file.write(script)
            script_path = Path(temp_file.name)
        completed = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script_path),
                str(path),
                TARGET_SHEET_NAME,
            ],
            check=False,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=120,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except FileNotFoundError:
        return SheetRename(path=path, success=False, message="no se encontro PowerShell para procesar archivos .xls")
    except subprocess.TimeoutExpired:
        return SheetRename(path=path, success=False, message="Excel no respondio al procesar el archivo .xls")
    finally:
        if script_path is not None:
            script_path.unlink(missing_ok=True)

    output = completed.stdout.strip()
    try:
        data = json.loads(output.splitlines()[-1] if output else "{}")
    except json.JSONDecodeError:
        data = {}
    if completed.returncode != 0 or not data.get("success"):
        message = str(data.get("message") or completed.stderr.strip() or "no se pudo procesar el archivo .xls con Excel")
        return SheetRename(path=path, success=False, message=message)
    return SheetRename(
        path=path,
        success=True,
        before=str(data.get("before") or ""),
        changed=bool(data.get("changed")),
    )
